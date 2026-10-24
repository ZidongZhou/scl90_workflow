#!/usr/bin/env python3
"""Run the reproducible analysis pipeline.

Default route:
python code/run_all.py --data /path/to/Data_collegestudent.csv --out outputs_final \
  --random-repeats 1000 --bootstrap-ci 500

The default route reruns the main models, calibration, top-k tables, selective
full-form escalation, domain-balanced item-set analyses, nested validation,
response-quality audits, and Figures 1-5. Random-item baseline and LASSO
selection-stability results are taken from the archived files in code/static/.

Full random-baseline recomputation:
python code/run_all.py --data /path/to/Data_collegestudent.csv --out outputs_full \
  --random-repeats 1000 --bootstrap-ci 500 --recompute-random-baseline

Tree-based sensitivity is optional and can be enabled with --include-tree or run
separately with code/tree_sensitivity.py.
"""
from __future__ import annotations
import argparse, csv, json, math, os, re, shutil, subprocess, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


def thread_limited_env() -> dict[str, str]:
    env = os.environ.copy()
    for key in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
        env.setdefault(key, "1")
    return env


def run(cmd: list[str]) -> None:
    print("\n$ " + " ".join(cmd), flush=True)
    subprocess.check_call(cmd, env=thread_limited_env())




def _csv_value(value: str):
    """Convert CSV numeric fields to real spreadsheet numbers; keep identifiers as text."""
    if value is None or value == "":
        return None
    if re.fullmatch(r"[+-]?\d+", value):
        digits = value.lstrip("+-")
        if len(digits) > 1 and digits.startswith("0"):
            return value
        return int(value)
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    if not math.isfinite(number):
        return value
    return number


def _read_csv_rows(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def _write_s1_table_workbook(summary_csv: Path, folds_csv: Path, out_xlsx: Path) -> None:
    """Create S1 Table with numeric cells, a summary sheet, and fold-level results."""
    workbook = Workbook()
    for index, (sheet_name, source) in enumerate(
        (("summary", summary_csv), ("folds", folds_csv))
    ):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = sheet_name
        for row_index, row in enumerate(_read_csv_rows(source)):
            worksheet.append(row if row_index == 0 else [_csv_value(cell) for cell in row])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in worksheet.columns:
            values = [len(str(cell.value)) if cell.value is not None else 0 for cell in column]
            worksheet.column_dimensions[column[0].column_letter].width = min(40, max(10, max(values, default=10) + 2))
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_xlsx)


SAFE_OUTPUT_NAMES = {"outputs_run", "outputs_final", "outputs_quick", "outputs_full", "outputs_smoke"}

def _is_dedicated_output_name(name: str) -> bool:
    return name in SAFE_OUTPUT_NAMES or any(
        name.startswith(prefix + "_")
        for prefix in ("outputs_run", "outputs_final", "outputs_quick", "outputs_full", "outputs_smoke")
    )

def validate_output_target(out: Path, package_root: Path, data: Path) -> Path:
    """Validate an output path before creating or removing anything."""
    if out.is_symlink():
        raise ValueError("Refusing an output path that is a symbolic link; choose a dedicated output directory.")
    target = out.resolve()
    root = package_root.resolve()
    if target == root or target in root.parents:
        raise ValueError(f"Refusing output path {target}: it is the project root or one of its parents.")
    try:
        relative = target.relative_to(root)
    except ValueError:
        relative = None
    if relative is not None:
        if len(relative.parts) != 1 or not _is_dedicated_output_name(relative.name):
            raise ValueError(
                f"Refusing output path {target}: within the project, use a dedicated top-level output directory "
                "such as outputs_run or outputs_final. Protected package directories cannot be used as output."
            )
    if target.exists() and not target.is_dir():
        raise ValueError(f"Output path exists and is not a directory: {target}")
    data_path = data.resolve()
    if target == data_path or target in data_path.parents:
        raise ValueError(f"Refusing output path {target}: it contains or is the input data file {data_path}.")
    if not data_path.exists() or not data_path.is_file():
        raise ValueError(f"Input data file does not exist or is not a regular file: {data_path}")
    return target

def validate_clean_target(out: Path, package_root: Path, data: Path) -> Path:
    """Only clean dedicated output directories, after validating the input data path."""
    target = validate_output_target(out, package_root, data)
    if not _is_dedicated_output_name(target.name):
        raise ValueError(
            f"Refusing --clean on {target}: use outputs_run, outputs_final, outputs_quick, outputs_full, outputs_smoke, or a suffixed variant."
        )
    return target

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True, type=Path, help="Path to Data_collegestudent.csv from the public Mendeley dataset")
    ap.add_argument("--out", default=Path("outputs_final"), type=Path)
    ap.add_argument("--random-repeats", default=1000, type=int)
    ap.add_argument("--bootstrap-ci", default=500, type=int)
    ap.add_argument("--recompute-random-baseline", action="store_true", help="Recompute random raw-score item-set baselines instead of copying archived static simulation outputs")
    ap.add_argument("--include-tree", action="store_true", help="Run the optional LightGBM/XGBoost item-ranking sensitivity analyses.")
    ap.add_argument("--quick", action="store_true", help="Quick check only; skips bootstrap resampling and nested cross-validation.")
    ap.add_argument("--skip-nested-cv", action="store_true", help="Skip nested cross-validation for a faster, non-final run.")
    ap.add_argument("--si-out", type=Path, default=None, help="Directory for SI-ready table copies. Defaults to <out>/supporting_info.")
    ap.add_argument("--clean", action="store_true", help="Remove an existing dedicated outputs_* directory after validating the input path")
    args = ap.parse_args()
    here = Path(__file__).resolve().parent
    package_root = here.parent
    if args.quick:
        args.bootstrap_ci = 0
        args.skip_nested_cv = True
        if args.recompute_random_baseline:
            args.random_repeats = min(args.random_repeats, 5)
    try:
        args.out = (
            validate_clean_target(args.out, package_root, args.data)
            if args.clean
            else validate_output_target(args.out, package_root, args.data)
        )
    except ValueError as exc:
        ap.error(str(exc))
    si_dir = args.si_out.resolve() if args.si_out is not None else args.out / "supporting_info"
    root_resolved = package_root.resolve()
    try:
        si_inside_root = si_dir.relative_to(root_resolved)
    except ValueError:
        si_inside_root = None
    if si_inside_root is not None and not (si_dir == args.out or args.out in si_dir.parents):
        ap.error("--si-out cannot write into the packaged source tree; choose a directory outside the project or under --out.")
    if si_dir == args.out or si_dir in args.out.parents:
        ap.error("--si-out must be a separate directory, not --out itself or one of its parent directories.")
    if args.clean and args.out.exists():
        shutil.rmtree(args.out)
    elif args.out.exists() and any(args.out.iterdir()):
        ap.error(
            f"Output directory is not empty: {args.out}. Use a new directory, or pass --clean "
            "for a dedicated outputs_* directory after checking its contents."
        )
    args.out.mkdir(parents=True, exist_ok=True)
    main_cmd = [sys.executable, str(here/"analysis.py"), "--data", str(args.data), "--out", str(args.out), "--random-repeats", str(args.random_repeats), "--bootstrap-ci", str(args.bootstrap_ci)]
    if args.recompute_random_baseline:
        main_cmd.append("--recompute-random-baseline")
    if args.quick:
        main_cmd.append("--quick")
    if args.skip_nested_cv:
        main_cmd.append("--skip-nested-cv")
    run(main_cmd)
    if args.include_tree:
        run([sys.executable, str(here/"tree_sensitivity.py"), "--data", str(args.data), "--out", str(args.out)])
    else:
        print("\nSkipping supplementary tree-based sensitivity by default. Use --include-tree to run it.", flush=True)
    run([sys.executable, str(here/"make_figures.py"), "--out", str(args.out), "--upload-dir", str(args.out / "figure_files_for_upload")])

    # Create SI-ready copies for the main supplementary tables. This keeps the
    # rerun package aligned with the manuscript Supporting information captions.
    si_dir.mkdir(parents=True, exist_ok=True)
    summary_csv = args.out / "tables" / "pooled_region_stratified_cv_summary.csv"
    folds_csv = args.out / "tables" / "pooled_region_stratified_cv_folds.csv"
    if summary_csv.exists() and folds_csv.exists():
        _write_s1_table_workbook(summary_csv, folds_csv, si_dir / "S1_Table.xlsx")

    table_map = {
        "lasso_bootstrap_stability_1000.csv": "S2_Table.csv",
        "tableS_tree_based_selection_sensitivity.csv": "S3_Table.csv",
        "random_rawsum_item_set_distributions.csv": "S4_Table.csv",
        "tableS_current_environment_lasso_sensitivity.csv": "S5_Table.csv",
        "tableS_data_quality_exclusion_sensitivity.csv": "S6_Table.csv",
        "tableS_topk_triage_bootstrap_ci.csv": "S7_Table.csv",
        "table11_selective_full_form_escalation.csv": "S8_Table.csv",
        "tableS_selective_escalation_bootstrap_ci.csv": "S9_Table.csv",
        "tableS_order_assumed_domain_balanced_item_selection.csv": "S10_Table.csv",
        "table12_nonoverlap_and_severity_validation.csv": "S11_Table.csv",
        "item_selection_cv_summary.csv": "S12_Table.csv",
        "item_selection_stability.csv": "S13_Table.csv",
        "table14_response_pattern_audit.csv": "S14_Table.csv",
        "tableS_response_quality_restricted_selective_escalation.csv": "S15_Table.csv",
        "tableS_nested_selective_escalation_cv_summary.csv": "S16_Table.csv",
        "tableS_selective_escalation_error_profile.csv": "S17_Table.csv",
        "tableS_subgroup_workflow_diagnostics.csv": "S18_Table.csv",
        "tableS_tripod_ai_informed_reporting_checklist.csv": "S19_Table.csv",
        "tableS_intended_use_deployment_card.csv": "S20_Table.csv",
        "tableS_selective_escalation_threshold_tradeoff.csv": "S21_Table.csv",
        "tableS_retrospective_response_quality_audit.csv": "S22_Table.csv",
        "tableS_first_stage_observable_quality_flags.csv": "S23_Table.csv",
    }
    for src_name, dst_name in table_map.items():
        src = args.out / "tables" / src_name
        archived = package_root / "supporting_info" / dst_name
        if dst_name == "S3_Table.csv" and not args.include_tree:
            if archived.exists():
                shutil.copy2(archived, si_dir / dst_name)
            continue
        if src.exists():
            shutil.copy2(src, si_dir / dst_name)
            continue
        if archived.exists():
            shutil.copy2(archived, si_dir / dst_name)

    config_path = args.out / "analysis_run_config.json"
    if config_path.exists():
        with config_path.open(encoding="utf-8") as f:
            config = json.load(f)
        config.update({
            "include_tree_requested": bool(args.include_tree),
            "tree_based_sensitivity_source": (
                "recomputed_current_run" if args.include_tree
                else "archived_packaged_output" if (si_dir / "S3_Table.csv").exists()
                else "not_available"
            ),
            "figure_export_formats": ["600-dpi LZW TIFF for submission", "PNG preview", "editable-text SVG", "editable-text PDF"],
            "wrapper_completed": True,
            "run_mode": (
                "quick" if args.quick else
                "custom" if args.skip_nested_cv or args.bootstrap_ci != 500 else
                "full"
            ),
            "full_nested_cv_run": not args.skip_nested_cv,
            "bootstrap_ci_replicates": args.bootstrap_ci,
        })
        config["package_version"] = "1.0.0"
        with config_path.open("w", encoding="utf-8") as f:
            json.dump(config, f, ensure_ascii=False, indent=2)
    environment_path = args.out / "environment_report.json"
    if environment_path.exists():
        with environment_path.open(encoding="utf-8") as f:
            environment = json.load(f)
        environment.update({
            "package_version": "1.0.0",
            "report_scope": "Runtime environment observed for this run; this report is not a dependency lock file.",
            "run_mode": (
                "quick" if args.quick else
                "custom" if args.skip_nested_cv or args.bootstrap_ci != 500 else
                "full"
            ),
            "tree_based_sensitivity_run": bool(args.include_tree),
        })
        with environment_path.open("w", encoding="utf-8") as f:
            json.dump(environment, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    main()
