"""Package and regression tests for the SCL-90 reproducibility workflow.

Run without raw data to verify package structure. To also run a lightweight raw
data check, set SCL90_TEST_DATA=/path/to/Data_collegestudent.csv before pytest.
The raw-data test intentionally uses the quick route and does not serve as the
submitted full recomputation route.
"""
from __future__ import annotations
import os, subprocess, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def test_required_static_and_si_files_exist():
    required = [
        ROOT / "code" / "run_all.py",
        ROOT / "code" / "analysis.py",
        ROOT / "code" / "static" / "lasso_rank_hunan.csv",
        ROOT / "supporting_info" / "S1_Table.xlsx",
        ROOT / "supporting_info" / "S6_Table.csv",
        ROOT / "supporting_info" / "S7_Table.csv",
        ROOT / "supporting_info" / "S8_Table.csv",
        ROOT / "supporting_info" / "S11_Table.csv",
        ROOT / "supporting_info" / "S12_Table.csv",
        ROOT / "LICENSE",
        ROOT / "CITATION.cff",
        ROOT / "requirements.txt",
        ROOT / "environment.yml",
    ]
    missing = [str(p.relative_to(ROOT)) for p in required if not p.exists()]
    assert not missing, f"Missing required package files: {missing}"


def test_all_supporting_tables_exist():
    for i in range(1, 24):
        matches = list((ROOT / "supporting_info").glob(f"S{i}_Table*"))
        assert matches, f"Missing S{i} Table"


def test_packaged_s1_workbook_opens_and_stores_metrics_as_numbers():
    import zipfile
    from openpyxl import load_workbook
    path = ROOT / "supporting_info" / "S1_Table.xlsx"
    with zipfile.ZipFile(path) as book:
        assert book.testzip() is None
        xml = book.read("xl/workbook.xml").decode("utf-8")
        assert 'name="summary"' in xml and 'name="folds"' in xml
    workbook = load_workbook(path, read_only=True, data_only=True)
    summary = workbook["summary"]
    metric_rows = {summary.cell(row, 1).value: row for row in range(2, summary.max_row + 1)}
    for metric in ("prevalence", "roc_auc", "pr_auc", "brier"):
        row = metric_rows[metric]
        assert isinstance(summary.cell(row, 2).value, (int, float)), metric
    workbook.close()


def test_no_obsolete_user_facing_terms_in_outputs():
    banned = [
        "Quality-gate-first",
        "response_pattern_audit_first",
        "Diagnostic subgroup only",
        "quality-gated selective escalation",
        "Domain-order-balanced domain-order-balanced",
    ]
    paths = list((ROOT / "supporting_info").glob("*")) + [
        ROOT / "README.md",
        ROOT / "MANIFEST.md",
        ROOT / "outputs" / "analysis_run_config.json",
    ]
    for path in paths:
        if path.suffix.lower() == ".xlsx" or not path.exists():
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        for term in banned:
            assert term not in text, f"{term} found in {path}"


def test_figure3_no_obsolete_response_pattern_audit_text():
    svg = ROOT / "outputs" / "figures" / "figure3_selective_escalation_workflow.svg"
    if not svg.exists():
        return
    text = svg.read_text(encoding="utf-8", errors="ignore")
    assert "Response-pattern audit" not in text
    assert "pattern checks" not in text


def test_run_all_help_exposes_recompute_and_include_tree():
    result = subprocess.run([sys.executable, str(ROOT / "code" / "run_all.py"), "--help"], text=True, capture_output=True, check=True)
    assert "--recompute-random-baseline" in result.stdout
    assert "--include-tree" in result.stdout
    assert "--skip-tree" not in result.stdout


def test_optional_minimal_pipeline_if_raw_data_available(tmp_path):
    data = os.environ.get("SCL90_TEST_DATA")
    if not data:
        return
    out = tmp_path / "outputs_smoke"
    subprocess.check_call([
        sys.executable, str(ROOT / "code" / "run_all.py"),
        "--data", data,
        "--out", str(out),
        "--clean",
        "--quick",
    ])
    for rel in [
        "tables/table1_sample_characteristics_quality.csv",
        "tables/table3_performance_by_item_count_and_selection.csv",
        "tables/table11_selective_full_form_escalation.csv",
        "tables/tableS_order_assumed_domain_balanced_item_selection.csv",
        "tables/tableS_response_quality_restricted_selective_escalation.csv",
        "tables/tableS_retrospective_response_quality_audit.csv",
        "tables/tableS_first_stage_observable_quality_flags.csv",
        "figures/figure1_burden_performance_frontier.png",
        "figure_files_for_upload/Fig1.tif",
        "figures/figure3_selective_escalation_workflow.png",
        "supporting_info/S1_Table.xlsx",
    ]:
        assert (out / rel).exists(), rel
    import zipfile
    with zipfile.ZipFile(out / "supporting_info" / "S1_Table.xlsx") as book:
        assert book.testzip() is None
        xml = book.read("xl/workbook.xml").decode("utf-8")
        assert 'name="summary"' in xml and 'name="folds"' in xml


def test_no_figure_numbers_above_five_are_packaged():
    import re
    for folder in [ROOT / "outputs" / "figures", ROOT / "figure_files_for_upload"]:
        if not folder.exists():
            continue
        for path in folder.iterdir():
            match = re.search(r"(?:figure|fig)(\d+)", path.stem, flags=re.IGNORECASE)
            if match:
                assert 1 <= int(match.group(1)) <= 5, f"Unexpected obsolete figure-numbered artifact: {path}"


def test_upload_tiffs_are_rgb_without_alpha():
    from PIL import Image

    for path in sorted((ROOT / "figure_files_for_upload").glob("Fig*.tif")):
        with Image.open(path) as image:
            assert image.mode in {"RGB", "L"}, f"Alpha channel found in {path}: {image.mode}"
            assert tuple(round(value) for value in image.info.get("dpi", ())) == (600, 600)


def test_quality_sensitivity_function_not_self_recursive():
    """Regression test for the accidentally shadowed function in the original archive."""
    import ast
    source = (ROOT / "code" / "analysis.py").read_text(encoding="utf-8")
    module = ast.parse(source)
    definitions = [n for n in module.body if isinstance(n, ast.FunctionDef) and n.name == "response_quality_restricted_selective_escalation_sensitivity"]
    assert len(definitions) == 1, "Duplicate definitions shadow the real sensitivity implementation"


def test_journal_si_names_are_short_and_complete():
    folder = ROOT / "supporting_info"
    assert sorted(p.name for p in folder.iterdir()) == sorted(
        ["S1_Table.xlsx"] + [f"S{i}_Table.csv" for i in range(2,24)]
    )


def test_development_thresholds_never_use_silent_hardcoded_fallback():
    code = (ROOT / "code" / "analysis.py").read_text(encoding="utf-8")
    assert 'raise RuntimeError("No development threshold pair meets' in code


def test_figure3_svg_text_is_at_least_8pt():
    import re
    txt = (ROOT / "outputs" / "figures" / "figure3_selective_escalation_workflow.svg").read_text(encoding="utf-8")
    sizes = [float(x) for x in re.findall(r"font-size: ([0-9.]+)px", txt)]
    # Matplotlib SVG uses CSS point sizes, encoded as px in the styling.
    assert not sizes or min(sizes) >= 8.0


def test_one_upload_format_and_no_duplicate_figure_folder():
    upload = ROOT / "figure_files_for_upload"
    assert sorted(p.name for p in upload.iterdir()) == [f"Fig{i}.tif" for i in range(1, 6)]
    assert not (ROOT / "outputs" / "figure_files_for_upload").exists()
    for i in range(1, 6):
        stem = [p for p in (ROOT / "outputs" / "figures").glob(f"figure{i}_*")]
        assert any(p.suffix == ".png" for p in stem), f"Missing PNG preview for Fig{i}"
        assert any(p.suffix == ".svg" for p in stem), f"Missing SVG source for Fig{i}"
        assert any(p.suffix == ".pdf" for p in stem), f"Missing PDF source for Fig{i}"
        assert not any(p.suffix == ".tif" for p in stem), f"Duplicate TIFF in outputs/figures for Fig{i}"


def test_font_report_is_explicit_without_machine_specific_paths():
    import json
    report = json.loads((ROOT / "outputs" / "figure_font_report.json").read_text(encoding="utf-8"))
    assert report["font_family_used"]
    assert report["font_file_name"]
    assert "font_file_used" not in report
    assert isinstance(report["true_arial_available"], bool)
    if not report["true_arial_available"]:
        assert report["font_family_used"].casefold() != "arial"


def test_raw_data_and_individual_outputs_are_gitignored():
    rules = (ROOT / ".gitignore").read_text(encoding="utf-8")
    assert "data_raw/" in rules
    assert "**/Data_collegestudent.csv" in rules
    assert "**/shandong_external_predictions_main_lasso15.csv" in rules
    assert "**/excluded_invalid_record.csv" in rules
    assert not (ROOT / "data_raw" / "Data_collegestudent.csv").exists()


def test_clean_refuses_to_remove_project_root():
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(ROOT / "not_a_real_input.csv"), "--out", str(ROOT), "--clean"],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "Refusing" in result.stderr
    assert (ROOT / "README.md").exists()


def test_clean_refuses_directory_containing_input_data(tmp_path):
    data = tmp_path / "Data_collegestudent.csv"
    data.write_text("placeholder", encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(data), "--out", str(tmp_path), "--clean"],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "input data file" in result.stderr
    assert data.exists()


def test_duplicate_random_baseline_output_aliases_are_not_packaged():
    tables = ROOT / "outputs" / "tables"
    assert not (tables / "random_rawsum_item_set_distributions_1000_each_size.csv").exists()
    assert not (tables / "table5_random_rawsum_item_set_summary_1000_each_size.csv").exists()
    assert not (ROOT / "requirements-lock.txt").exists()


def test_release_version_metadata_is_consistent():
    import json
    import re
    citation = (ROOT / "CITATION.cff").read_text(encoding="utf-8")
    config = json.loads((ROOT / "outputs" / "analysis_run_config.json").read_text(encoding="utf-8"))
    changelog = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
    assert re.search(r'^version: "1\.0\.0"$', citation, flags=re.MULTILINE)
    assert config["package_version"] == "1.0.0"
    assert re.search(r"^## 1\.0\.0$", changelog, flags=re.MULTILINE)


def test_clean_refuses_protected_project_subdirectories():
    for folder in ("code", "outputs", "supporting_info", "tests"):
        result = subprocess.run(
            [sys.executable, str(ROOT / "code" / "run_all.py"),
             "--data", str(ROOT / "missing_input.csv"), "--out", str(ROOT / folder), "--clean"],
            text=True, capture_output=True,
        )
        assert result.returncode != 0, folder
        assert "Refusing" in result.stderr, (folder, result.stderr)
    assert (ROOT / "code" / "run_all.py").exists()
    assert (ROOT / "outputs" / "analysis_run_config.json").exists()
    assert (ROOT / "supporting_info" / "S1_Table.xlsx").exists()

def test_clean_validates_input_before_removing_existing_output(tmp_path):
    out = tmp_path / "outputs_smoke"
    out.mkdir()
    sentinel = out / "keep.txt"
    sentinel.write_text("keep", encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(tmp_path / "missing.csv"), "--out", str(out), "--clean"],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "Input data file does not exist" in result.stderr
    assert sentinel.read_text(encoding="utf-8") == "keep"

def test_clean_refuses_non_dedicated_directory_names(tmp_path):
    out = tmp_path / "important_files"
    out.mkdir()
    data = tmp_path / "Data_collegestudent.csv"
    data.write_text("placeholder", encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(data), "--out", str(out), "--clean"],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "use outputs_run" in result.stderr

def test_internal_script_and_archive_names_are_compact():
    for name in ("analysis.py", "make_figures.py", "tree_sensitivity.py", "check_font.py"):
        assert (ROOT / "code" / name).is_file(), name
    static_names = {p.name for p in (ROOT / "code" / "static").iterdir()}
    assert static_names == {
        "random_sets_1000.csv", "random_summary_1000.csv", "lasso_rank_hunan.csv",
        "lasso_rank_shandong.csv", "lasso_stability_1000.csv"
    }
    tables = ROOT / "outputs" / "tables"
    assert (tables / "item_selection_cv_summary.csv").exists()
    assert (tables / "item_selection_cv_folds.csv").exists()
    assert (tables / "item_selection_stability.csv").exists()
    assert not list(tables.glob("table13_fully_nested_item_selection_*.csv"))


def test_nonempty_output_requires_clean_or_a_new_directory(tmp_path):
    out = tmp_path / "outputs_smoke"
    out.mkdir()
    sentinel = out / "old_result.csv"
    sentinel.write_text("old", encoding="utf-8")
    data = tmp_path / "Data_collegestudent.csv"
    data.write_text("placeholder", encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(data), "--out", str(out)],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "Output directory is not empty" in result.stderr
    assert sentinel.read_text(encoding="utf-8") == "old"


def test_clean_refuses_output_symlink(tmp_path):
    data = tmp_path / "Data_collegestudent.csv"
    data.write_text("placeholder", encoding="utf-8")
    target = tmp_path / "real_output"
    target.mkdir()
    sentinel = target / "keep.txt"
    sentinel.write_text("keep", encoding="utf-8")
    link = tmp_path / "outputs_smoke"
    try:
        link.symlink_to(target, target_is_directory=True)
    except (OSError, NotImplementedError):
        return
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(data), "--out", str(link), "--clean"],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "symbolic link" in result.stderr
    assert sentinel.read_text(encoding="utf-8") == "keep"

def test_si_output_cannot_target_packaged_supporting_info_before_clean(tmp_path):
    out = tmp_path / "outputs_smoke"
    out.mkdir()
    sentinel = out / "keep.txt"
    sentinel.write_text("keep", encoding="utf-8")
    data = tmp_path / "Data_collegestudent.csv"
    data.write_text("placeholder", encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(ROOT / "code" / "run_all.py"),
         "--data", str(data), "--out", str(out), "--si-out", str(ROOT / "supporting_info"), "--clean"],
        text=True, capture_output=True,
    )
    assert result.returncode != 0
    assert "packaged source tree" in result.stderr
    assert sentinel.read_text(encoding="utf-8") == "keep"
